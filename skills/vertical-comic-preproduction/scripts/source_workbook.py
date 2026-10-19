"""Read an XLSX as raw cell evidence, without interpreting panel IDs or approving it."""
from pathlib import Path
import argparse
import datetime
import hashlib
import json
from platform_io import configure_stdio


def snapshot(path):
    from openpyxl import load_workbook
    path = Path(path).resolve()
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    sheets = []
    # Explicit binary stream guarantees closure on Windows, including read errors.
    with path.open('rb') as stream:
        book = load_workbook(stream, read_only=True, data_only=False)
        try:
            for sheet in book:
                rows = []
                for row in sheet.iter_rows():
                    cells = []
                    for cell in row:
                        if cell.value is None:
                            continue
                        value = cell.value
                        if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
                            value = value.isoformat()
                        cells.append({'coordinate': cell.coordinate, 'type': cell.data_type, 'value': value})
                    if cells:
                        rows.append(cells)
                sheets.append({'name': sheet.title, 'rows': rows})
        finally:
            book.close()
    if hashlib.sha256(path.read_bytes()).hexdigest() != before:
        raise ValueError('源工作簿读取期间变化，请重新读取')
    return {'policy': 'raw_workbook_evidence_v1', 'source_sha256': before,
            'sheets': sheets, 'status': 'RAW_EVIDENCE_NOT_SEMANTIC_APPROVAL'}


if __name__ == '__main__':
    configure_stdio()
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', required=True, type=Path)
    parser.add_argument('--output', required=True, type=Path)
    args = parser.parse_args()
    result = snapshot(args.source)
    with args.output.open('x', encoding='utf-8', newline='\n') as stream:
        json.dump(result, stream, ensure_ascii=False, indent=2)
        stream.write('\n')
    print(json.dumps({'status': result['status'], 'output': str(args.output.resolve())}, ensure_ascii=False))
